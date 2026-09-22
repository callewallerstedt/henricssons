#!/usr/bin/env python3
"""Importerar leverantörens lista med hamnkapell och överdrag (havnekalesje.no).

Läser antingen leverantörens Excel-fil direkt eller en redan parsad
artiklar.json, speglar hem bilderna till henricssons_bilder/leverantor/,
skriver supplier_products_seed.json och upsertar artiklarna i databasen.

    python import_supplier_products.py Havnekalesje_modellista.xlsx            # torrkörning
    python import_supplier_products.py Havnekalesje_modellista.xlsx --commit   # skriver
    python import_supplier_products.py supplier_products_seed.json              # synka mot seed-filen

Utan --commit skrivs ingenting: skriptet visar bara vad som skulle ändras.

Upserten matchar på slug. Pris, artikelnummer, årsmodell och bild uppdateras
från filen. Modell, variant och beskrivning sätts bara när artikeln skapas, så
att texter som svenskats i adminpanelen inte skrivs över av leverantörens
norska. Leveranstid, publicering och sortering rörs aldrig.

Databasen väljs som i appen: DATABASE_URL i miljön (eller .env), annars den
lokala SQLite-filen.
"""

import argparse
import json
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

BASE_DIR = Path(__file__).resolve().parent
SEED_FILE = BASE_DIR / "supplier_products_seed.json"
SV_FILE = BASE_DIR / "supplier_products_sv.json"
SV_FIELDS = ("model", "model_year", "variant", "description")
IMAGES_ROOT = BASE_DIR / "henricssons_bilder"
IMAGE_DIR = "leverantor"
IMAGE_MAX_EDGE = 1600

# Fliknamn i leverantörens Excel -> (produkttyp, varumärke)
SHEETS = {
    "Yamarin hamnkapell": ("Hamnkapell", "Yamarin"),
    "Yamarin konsollöverdrag": ("Konsollöverdrag", "Yamarin"),
    "Buster hamnkapell": ("Hamnkapell", "Buster"),
    "Buster stolöverdrag": ("Stolöverdrag", "Buster"),
}

# Artiklar som leverantören själv hänvisar vidare från. De visas men går inte
# att beställa; detaljsidan länkar till ersättaren.
REPLACEMENTS = {
    "YAM49BRCR": "YAM50BR",
    "BUSXLVMAX": "BUSXL1",
}


def slugify(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()


def brand_of(model: str, base_brand: str) -> str:
    # Excel-filen har ingen egen Cross-flik; Cross känns igen på modellnamnet.
    if base_brand == "Yamarin" and re.search(r"\bcross\b", model, re.I):
        return "Yamarin Cross"
    return base_brand


def parse_excel(path: Path) -> List[Dict[str, Any]]:
    import openpyxl  # bara nödvändig när en ny Excel-fil läses in

    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for sheet_name, (product_type, base_brand) in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            print(f"VARNING: fliken {sheet_name!r} saknas i filen")
            continue
        for raw in wb[sheet_name].iter_rows(min_row=3, values_only=True):
            model = str(raw[0] or "").strip()
            if not model:
                continue
            gross = raw[7]
            rows.append({
                "artikelnummer": str(raw[1] or "").strip(),
                "modell": model,
                "varumarke": brand_of(model, base_brand),
                "produkttyp": product_type,
                "arsmodell": str(raw[2] or "").strip(),
                "variant": str(raw[3] or "").strip(),
                "beskrivning": str(raw[4] or "").strip(),
                "bildlank": str(raw[5] or "").strip(),
                "kalla": str(raw[6] or "").strip(),
                "pris_brutto_raw": float(gross) if isinstance(gross, (int, float)) else 0.0,
            })
    return rows


def load_source(path: Path) -> List[Dict[str, Any]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return parse_excel(path)
    return json.loads(path.read_text(encoding="utf-8"))


def local_image_path(url: str) -> str:
    """https://havnekalesje.no/media/pages/produkter/<märke>/<typ>/<modell>/<hash>/<fil>
    -> leverantor/<märke>/<modell>-<fil>. Hashen hoppas över eftersom den ändras
    varje gång leverantören bygger om sin sajt."""
    parts = [p for p in urlparse(url).path.split("/") if p]
    filename = parts[-1] if parts else "bild.jpg"
    stem, dot, ext = filename.rpartition(".")
    filename = f"{slugify(stem or filename)}.{ext.lower()}" if dot else slugify(filename)
    brand = slugify(parts[3]) if len(parts) > 3 else "okand"
    model_folder = slugify(parts[-3]) if len(parts) >= 3 else ""
    return f"{IMAGE_DIR}/{brand}/{model_folder}-{filename}" if model_folder else f"{IMAGE_DIR}/{brand}/{filename}"


def normalize_rows(source_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    products: List[Dict[str, Any]] = []
    seen_slugs: Dict[str, int] = {}
    seen_articles: Dict[str, str] = {}
    for index, row in enumerate(source_rows):
        article_no = str(row.get("artikelnummer") or "").strip()
        model = str(row.get("modell") or "").strip()
        brand = str(row.get("varumarke") or "").strip()
        price = float(row.get("pris_brutto_raw") or 0)

        slug = slugify(article_no or model)
        if slug in seen_slugs:
            # YAM63BR1 finns på två rader i leverantörens fil. Båda behålls så
            # att Niclas kan reda ut vilken som är rätt.
            slug = f"{slug}-{slugify(model)}"
            print(f"VARNING: artikelnummer {article_no} används på flera rader, "
                  f"'{model}' får slug {slug}")
        seen_slugs[slug] = index
        if article_no.upper() in seen_articles and seen_articles[article_no.upper()] != model:
            print(f"VARNING: dubblerat artikelnummer {article_no}: "
                  f"'{seen_articles[article_no.upper()]}' och '{model}'")
        seen_articles.setdefault(article_no.upper(), model)

        # "Yamarin 46 SC & Cross" gäller både vanliga Yamarin och Cross.
        also_listed_under = "Yamarin" if brand == "Yamarin Cross" and "& cross" in model.lower() else ""
        replacement = REPLACEMENTS.get(article_no.upper(), "")
        image_url = str(row.get("bildlank") or "").strip()

        products.append({
            "slug": slug,
            "article_no": article_no,
            "model": model,
            "brand": brand,
            "also_listed_under": also_listed_under,
            "product_type": str(row.get("produkttyp") or "").strip(),
            "model_year": str(row.get("arsmodell") or "").strip(),
            "variant": str(row.get("variant") or "").strip(),
            "description": str(row.get("beskrivning") or "").strip(),
            "source_url": str(row.get("kalla") or "").strip(),
            "source_image_url": image_url,
            "image_path": local_image_path(image_url) if image_url else "",
            "price_ore": int(round(price * 100)),
            "orderable": bool(price) and not replacement,
            "replacement_article_no": replacement,
            "sort_order": index,
        })
    return products


def apply_swedish_texts(products: List[Dict[str, Any]]) -> None:
    """Leverantörens texter är på norska. supplier_products_sv.json har svenska
    texter per slug och läggs över vid varje import, så att en ny Excel-fil
    inte tar tillbaka norskan. Nya artiklar utan översättning listas."""
    translations = json.loads(SV_FILE.read_text(encoding="utf-8")) if SV_FILE.exists() else {}
    missing = []
    for product in products:
        texts = translations.get(product["slug"])
        if texts:
            product.update({k: v for k, v in texts.items() if k in SV_FIELDS})
        else:
            missing.append(product["slug"])
    if missing:
        print(f"Saknar svensk text (visas på norska tills {SV_FILE.name} fylls på): {', '.join(missing)}")


def shrink_image(data: bytes, suffix: str) -> bytes:
    """Leverantörens bilder är upp till 5 MB mobilfoton. Skala ned till samma
    storleksordning som resten av henricssons_bilder så att repot hålls lätt;
    servern gör ändå mindre webp-varianter vid behov."""
    try:
        from io import BytesIO
        from PIL import Image, ImageOps
    except ImportError:
        return data
    try:
        with Image.open(BytesIO(data)) as img:
            img = ImageOps.exif_transpose(img)
            img.thumbnail((IMAGE_MAX_EDGE, IMAGE_MAX_EDGE))
            out = BytesIO()
            if suffix.lower() == ".webp":
                img.save(out, "WEBP", quality=82)
            else:
                img.convert("RGB").save(out, "JPEG", quality=82, optimize=True, progressive=True)
            shrunk = out.getvalue()
        return shrunk if len(shrunk) < len(data) else data
    except Exception as exc:
        print(f"  kunde inte skala om bilden, sparar originalet: {exc}")
        return data


def mirror_images(products: List[Dict[str, Any]], commit: bool) -> None:
    wanted: Dict[str, str] = {}
    for product in products:
        if product["image_path"] and product["source_image_url"]:
            wanted.setdefault(product["image_path"], product["source_image_url"])
    missing = {rel: url for rel, url in wanted.items() if not (IMAGES_ROOT / rel).exists()}
    print(f"Bilder: {len(wanted)} unika, {len(missing)} saknas lokalt")
    if not commit:
        return
    for rel, url in sorted(missing.items()):
        target = IMAGES_ROOT / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        request = urllib.request.Request(url, headers={"User-Agent": "Henricssons-import/1.0"})
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                data = response.read()
            if not data:
                raise ValueError("tomt svar")
            target.write_bytes(shrink_image(data, target.suffix))
            print(f"  hämtad {rel} ({target.stat().st_size // 1024} KB)")
        except Exception as exc:
            print(f"  FEL: kunde inte hämta {url}: {exc}")


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", help="Leverantörens .xlsx eller en parsad artiklar.json")
    parser.add_argument("--commit", action="store_true", help="Skriv bilder, seed-fil och databas")
    parser.add_argument("--skip-db", action="store_true", help="Uppdatera bara bilder och seed-fil")
    args = parser.parse_args(argv)

    source_rows = load_source(Path(args.source))
    if source_rows and all("article_no" in row and "slug" in row for row in source_rows):
        # supplier_products_seed.json är redan normaliserad: synka databasen mot den.
        products = source_rows
    else:
        products = normalize_rows(source_rows)
    apply_swedish_texts(products)
    zero_price = [p["article_no"] for p in products if not p["price_ore"]]
    print(f"{len(products)} artiklar")
    if zero_price:
        print(f"Utan pris (ej direktbeställningsbara): {', '.join(zero_price)}")

    mirror_images(products, args.commit)

    if args.commit:
        SEED_FILE.write_text(json.dumps(products, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Skrev {SEED_FILE.name}")

    if args.skip_db:
        return 0

    import admin_api_flask as app_module

    report = app_module.upsert_supplier_products(products, commit=args.commit)
    if report is None:
        print("Databasen är inte nåbar, inget skrevs dit.")
        return 1
    print(f"Databas: {report['created']} nya, {report['updated']} uppdaterade, {report['unchanged']} oförändrade")
    for line in report["changes"]:
        print(f"  {line}")
    if not args.commit:
        print("\nTorrkörning - kör igen med --commit för att skriva.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
